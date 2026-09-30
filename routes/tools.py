import io
import logging
import mimetypes
import os
from pathlib import Path

from flask import Blueprint, jsonify, render_template, request, send_file, session

from services.ocr_service import ALLOWED_EXTENSIONS, MAX_FILE_BYTES, SUPPORTED_LANGUAGES, process_document

bp = Blueprint('tools', __name__)
logger = logging.getLogger(__name__)


def _auth():
    return session.get('user_id')


def _allowed_file(upload):
    if not upload or not upload.filename:
        return False
    ext = Path(upload.filename).suffix.lower().lstrip('.')
    return ext in ALLOWED_EXTENSIONS


def _valid_signature(upload):
    upload.stream.seek(0)
    head = upload.stream.read(16)
    upload.stream.seek(0)
    ext = Path(upload.filename).suffix.lower()
    signatures = {
        '.pdf': head.startswith(b'%PDF-'),
        '.png': head.startswith(b'\x89PNG\r\n\x1a\n'),
        '.jpg': head.startswith(b'\xff\xd8\xff'),
        '.jpeg': head.startswith(b'\xff\xd8\xff'),
        '.gif': head.startswith((b'GIF87a', b'GIF89a')),
        '.bmp': head.startswith(b'BM'),
        '.tif': head.startswith((b'II*\x00', b'MM\x00*')),
        '.tiff': head.startswith((b'II*\x00', b'MM\x00*')),
    }
    return signatures.get(ext, False)


@bp.get('/tools')
@bp.get('/tools/ocr')
def tools_ocr():
    if not _auth():
        return render_template('403.html'), 403
    return render_template('ocr.html', languages=SUPPORTED_LANGUAGES)


@bp.post('/api/tools/ocr')
def ocr():
    if not _auth():
        return jsonify(ok=False, message='No autenticado.'), 401
    upload = request.files.get('file')
    language = str(request.form.get('language', 'spanish')).strip().lower()
    if not upload:
        return jsonify(ok=False, message='Selecciona un archivo.'), 400
    if not _allowed_file(upload):
        return jsonify(ok=False, message='Formato no compatible. Usa PDF, PNG, JPG, JPEG, TIFF, BMP o GIF.'), 400

    upload.stream.seek(0, 2)
    size = upload.stream.tell()
    upload.stream.seek(0)
    if size > MAX_FILE_BYTES:
        return jsonify(ok=False, message='El archivo supera el máximo permitido por el servicio OCR (100 MB).'), 413
    if size == 0:
        return jsonify(ok=False, message='El archivo está vacío.'), 400
    if not _valid_signature(upload):
        return jsonify(ok=False, message='El contenido del archivo no coincide con su extensión.'), 400

    try:
        result = process_document(upload, language)
        if not result['text']:
            return jsonify(ok=False, message='El OCR terminó pero no detectó texto.'), 422
        return jsonify(ok=True, **result)
    except PermissionError as exc:
        return jsonify(ok=False, message=str(exc)), 502
    except (ValueError, ConnectionError, RuntimeError) as exc:
        logger.warning('OCR user-facing error: %s', exc)
        return jsonify(ok=False, message=str(exc)), 502
    except Exception:
        logger.exception('Unexpected OCR failure')
        return jsonify(ok=False, message='No se pudo procesar el documento. Intenta nuevamente.'), 502


@bp.post('/api/tools/ocr/export')
def ocr_export():
    if not _auth():
        return jsonify(ok=False, message='No autenticado.'), 401
    data = request.get_json() or {}
    text = str(data.get('text', ''))
    fmt = str(data.get('format', 'txt')).lower()
    if not text.strip():
        return jsonify(ok=False, message='No hay texto para exportar.'), 400
    if len(text) > 2_000_000:
        return jsonify(ok=False, message='El resultado es demasiado grande para esta exportación.'), 413

    if fmt == 'txt':
        return send_file(io.BytesIO(text.encode('utf-8')), mimetype='text/plain; charset=utf-8', as_attachment=True, download_name='ocr_resultado.txt')

    if fmt == 'docx':
        from docx import Document
        doc = Document()
        for paragraph in text.split('\n'):
            doc.add_paragraph(paragraph)
        stream = io.BytesIO(); doc.save(stream); stream.seek(0)
        return send_file(stream, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document', as_attachment=True, download_name='ocr_resultado.docx')

    if fmt == 'xlsx':
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = 'OCR'
        for row, line in enumerate(text.splitlines() or [''], 1):
            ws.cell(row=row, column=1, value=line)
        stream = io.BytesIO(); wb.save(stream); stream.seek(0)
        return send_file(stream, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='ocr_resultado.xlsx')

    if fmt == 'pdf':
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        stream = io.BytesIO(); pdf = canvas.Canvas(stream, pagesize=A4)
        width, height = A4; y = height - 45
        pdf.setFont('Helvetica', 10)
        for line in text.splitlines() or ['']:
            while len(line) > 110:
                pdf.drawString(40, y, line[:110]); line = line[110:]; y -= 14
                if y < 45: pdf.showPage(); pdf.setFont('Helvetica', 10); y = height - 45
            pdf.drawString(40, y, line); y -= 14
            if y < 45: pdf.showPage(); pdf.setFont('Helvetica', 10); y = height - 45
        pdf.save(); stream.seek(0)
        return send_file(stream, mimetype='application/pdf', as_attachment=True, download_name='ocr_resultado.pdf')

    return jsonify(ok=False, message='Formato de exportación no permitido.'), 400
